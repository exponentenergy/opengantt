import frappe
import json
import os

from opengantt.api.parse_upload import bulk_insert_tasks


@frappe.whitelist()
def reimport():
    """Re-parse the source_file attached to an OG Gantt against its template.

    Body: { "gantt": "<gantt name>" }. Returns { ok, count }.
    Replaces all existing tasks for that Gantt.
    """
    data = json.loads(frappe.request.data or "{}")
    gantt_name = data.get("gantt")
    if not gantt_name:
        frappe.throw("gantt is required")

    gantt = frappe.get_doc("OG Gantt", gantt_name)
    if not gantt.source_file:
        frappe.throw("No source_file attached to this Gantt — re-upload via New Gantt flow.")
    if not gantt.template:
        frappe.throw("Gantt has no template")

    template = frappe.get_doc("OG Template", gantt.template)
    field_map = json.loads(template.field_map or "{}")
    grouping = json.loads(template.grouping or "[]")

    rows = _read_source(gantt.source_file)
    tasks = _build_tasks(rows, field_map, grouping)
    count, _ = bulk_insert_tasks(gantt_name, tasks)
    frappe.db.set_value("OG Gantt", gantt_name, "parsed_at", frappe.utils.now_datetime())
    return {"ok": True, "count": count}


def _read_source(file_url):
    """Read an attached File doc as a list of dicts keyed by header."""
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    full_path = file_doc.get_full_path()
    ext = os.path.splitext(full_path)[1].lower()
    if ext in (".xlsx", ".xls"):
        return _read_xlsx(full_path)
    if ext in (".csv", ".tsv"):
        return _read_csv(full_path, delim="\t" if ext == ".tsv" else ",")
    frappe.throw(f"Unsupported file extension: {ext}")


def _read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(v is not None for v in r):
            continue
        out.append({headers[i]: r[i] for i in range(len(headers))})
    return out


def _read_csv(path, delim=","):
    import csv

    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=delim)
        return [r for r in reader]


def _build_tasks(rows, field_map, grouping):
    """Mirror of the frontend parser. Pure Python implementation."""
    tasks = []
    group_cache = {}  # path-tuple -> temp_id

    def _temp(prefix, i):
        return f"{prefix}-{i}"

    def _iso(v):
        if v is None or v == "":
            return None
        if hasattr(v, "isoformat"):
            try:
                return v.date().isoformat() if hasattr(v, "date") else v.isoformat()
            except Exception:
                return str(v)
        return str(v)

    name_col = field_map.get("name")
    start_col = field_map.get("start_date")
    end_col = field_map.get("end_date")
    a_start_col = field_map.get("actual_start")
    a_end_col = field_map.get("actual_end")

    for i, row in enumerate(rows):
        path = []
        parent_temp = None
        for g_col in grouping:
            val = row.get(g_col)
            if val is None or val == "":
                continue
            path.append(str(val))
            key = tuple(path)
            if key not in group_cache:
                tid = _temp("g", len(group_cache))
                group_cache[key] = tid
                tasks.append({
                    "temp_id": tid,
                    "parent_temp_id": parent_temp,
                    "task_name": str(val),
                    "kind": "group",
                    "sort_order": len(group_cache),
                    "fields": {},
                })
            parent_temp = group_cache[key]

        leaf_name = row.get(name_col) if name_col else None
        if leaf_name is None or leaf_name == "":
            continue
        fields = {k: v for k, v in row.items() if k not in (
            name_col, start_col, end_col, a_start_col, a_end_col, *grouping
        ) and v not in (None, "")}
        # coerce non-serialisable types to strings
        fields = {k: (v if isinstance(v, (str, int, float, bool)) else str(v)) for k, v in fields.items()}
        tasks.append({
            "temp_id": _temp("l", i),
            "parent_temp_id": parent_temp,
            "task_name": str(leaf_name),
            "kind": "leaf",
            "start_date": _iso(row.get(start_col)) if start_col else None,
            "end_date": _iso(row.get(end_col)) if end_col else None,
            "actual_start": _iso(row.get(a_start_col)) if a_start_col else None,
            "actual_end": _iso(row.get(a_end_col)) if a_end_col else None,
            "sort_order": i,
            "fields": fields,
        })
    return tasks

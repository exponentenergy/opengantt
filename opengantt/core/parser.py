"""Single server-side parser for OpenGantt imports.

Reads CSV/TSV (utf-8-sig then latin-1 fallback) and XLSX (openpyxl,
data_only, first sheet). Auto-detects the header row, coerces dates
column-wide (date/datetime objects, Excel serials, ISO, D/M/Y vs M/D/Y
disambiguated per-column, month-name formats) and offers fuzzy column
role detection for the Smart Import wizard.

Pure stdlib + openpyxl + dateutil. No frappe imports — usable from
install/seed code and API code alike.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import io
import os
import re
from datetime import date, datetime, timedelta

try:
    from dateutil import parser as du_parser
except Exception:  # pragma: no cover
    du_parser = None

SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".xlsx")

_EXCEL_EPOCH = datetime(1899, 12, 30)

# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


class ParserError(Exception):
    pass


def read_matrix(path: str) -> list[list]:
    """Read a CSV/TSV/XLSX file into a raw matrix of cell values."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return _read_xlsx(path)
    if ext in (".csv", ".tsv"):
        return _read_delimited(path, "\t" if ext == ".tsv" else ",")
    raise ParserError(
        f"Unsupported file type '{ext}'. Please upload a .csv, .tsv or .xlsx file."
    )


def _read_xlsx(path: str) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    matrix = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return matrix


def _read_delimited(path: str, delim: str) -> list[list]:
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover — latin-1 never fails
        text = raw.decode("latin-1", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [row for row in reader]


# ---------------------------------------------------------------------------
# Header detection
# ---------------------------------------------------------------------------


def detect_header_row(matrix: list[list]) -> int:
    """Pick the first row where >=60% of cells are non-empty and unique."""
    for idx, row in enumerate(matrix[:20]):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
        if not row or len(row) == 0:
            continue
        width = max(len(r) for r in matrix[: idx + 5] or [row])
        if width == 0:
            continue
        if len(cells) >= max(2, int(0.6 * len(row))) and len(set(cells)) == len(cells):
            return idx
    return 0


def _uniquify(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for h in headers:
        h = h or "Column"
        if h in seen:
            seen[h] += 1
            out.append(f"{h} ({seen[h]})")
        else:
            seen[h] = 1
            out.append(h)
    return out


# ---------------------------------------------------------------------------
# Date coercion
# ---------------------------------------------------------------------------

_ISO_RE = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}([T ].*)?$")
_SLASH_RE = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$")
_YMD_SLASH_RE = re.compile(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})$")
_MONTHNAME_RE = re.compile(
    r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)", re.IGNORECASE
)


def looks_like_date(value) -> bool:
    """Cheap check whether a single cell value could be a date."""
    if value is None or value == "":
        return False
    if isinstance(value, (datetime, date)):
        return True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return 20000 <= float(value) <= 80000
    s = str(value).strip()
    if not s:
        return False
    if _ISO_RE.match(s) or _SLASH_RE.match(s) or _YMD_SLASH_RE.match(s):
        return True
    if _MONTHNAME_RE.search(s) and any(ch.isdigit() for ch in s):
        return True
    return False


def coerce_date(value, dayfirst: bool = False) -> date | None:
    """Coerce one cell to a datetime.date, or None if not a date."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= float(value) <= 80000:
            return (_EXCEL_EPOCH + timedelta(days=float(value))).date()
        return None
    s = str(value).strip()
    if not s:
        return None
    m = _YMD_SLASH_RE.match(s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    if _ISO_RE.match(s):
        try:
            return date.fromisoformat(s[:10].replace("/", "-"))
        except ValueError:
            pass
    m = _SLASH_RE.match(s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        dd, mm = (a, b) if dayfirst else (b, a)
        if mm > 12 and dd <= 12:
            dd, mm = mm, dd
        try:
            return date(y, mm, dd)
        except ValueError:
            return None
    if _MONTHNAME_RE.search(s) and du_parser is not None:
        try:
            return du_parser.parse(s, dayfirst=dayfirst, fuzzy=False).date()
        except (ValueError, OverflowError):
            return None
    return None


def column_dayfirst(values: list) -> bool:
    """Column-wide D/M/Y vs M/D/Y disambiguation for slash-formatted cells."""
    saw_first_gt_12 = False
    saw_second_gt_12 = False
    for v in values:
        if v is None:
            continue
        m = _SLASH_RE.match(str(v).strip())
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a > 12:
            saw_first_gt_12 = True
        if b > 12:
            saw_second_gt_12 = True
    if saw_first_gt_12 and not saw_second_gt_12:
        return True
    return False


def date_ratio(values: list) -> float:
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return 0.0
    hits = sum(1 for v in non_empty if looks_like_date(v))
    return hits / len(non_empty)


# ---------------------------------------------------------------------------
# parse_file — the single entry point
# ---------------------------------------------------------------------------


def parse_file(path: str) -> dict:
    """Parse a spreadsheet into {headers, rows, date_columns, dayfirst}.

    rows are dicts keyed by (uniquified) header. Cells of date-like columns
    are coerced to datetime.date; everything else keeps its scalar value
    (strings stripped).
    """
    matrix = read_matrix(path)
    matrix = [row for row in matrix if row is not None]
    if not matrix:
        return {"headers": [], "rows": [], "date_columns": [], "dayfirst": {}}

    h_idx = detect_header_row(matrix)
    headers = _uniquify(
        [str(c).strip() if c is not None else "" for c in matrix[h_idx]]
    )
    body = matrix[h_idx + 1 :]

    # Column analysis for date coercion
    columns: dict[str, list] = {h: [] for h in headers}
    for row in body:
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
        for i, h in enumerate(headers):
            columns[h].append(row[i] if i < len(row) else None)

    date_columns = [h for h in headers if h and date_ratio(columns[h]) >= 0.6]
    dayfirst = {h: column_dayfirst(columns[h]) for h in date_columns}

    rows = []
    for row in body:
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            v = row[i] if i < len(row) else None
            if isinstance(v, str):
                v = v.strip()
            if h in date_columns:
                coerced = coerce_date(v, dayfirst=dayfirst.get(h, False))
                rec[h] = coerced if coerced is not None else (v or None)
            else:
                rec[h] = v if v not in ("",) else None
        rows.append(rec)

    return {
        "headers": [h for h in headers if h],
        "rows": rows,
        "date_columns": date_columns,
        "dayfirst": dayfirst,
    }


# ---------------------------------------------------------------------------
# Role detection
# ---------------------------------------------------------------------------

_ROLE_SYNONYMS = {
    "task_name": [
        "task name", "task", "name", "activity", "step", "item", "title",
        "program step", "project step", "module", "module name", "deliverable",
        "work item", "milestone name", "description",
    ],
    "start_date": ["start date", "start", "begin", "begins", "from", "planned start", "kickoff"],
    "end_date": [
        "end date", "end", "finish", "due", "due date", "to", "completion",
        "planned end", "deadline", "target date",
    ],
    "actual_start": ["actual start", "actual start date", "started on"],
    "actual_end": ["actual end", "actual end date", "actual finish", "completed on", "done on"],
    "owner": ["owner", "assignee", "assigned to", "responsible", "resource", "who", "lead"],
    "status": ["status", "state", "phase", "stage", "progress status"],
    "progress": [
        "progress", "% complete", "percent complete", "completion %", "% done",
        "done %", "percentage", "completion",
    ],
}

_GROUPING_PRIORITY = [
    "phase", "stage", "workstream", "section", "category", "group",
    "epic", "milestone", "sprint", "owner", "team", "status",
]


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9% ]+", " ", (s or "").lower()).strip()


def header_fingerprint(headers: list[str]) -> str:
    """Stable hash of the (sorted, normalized) header set for template matching."""
    canon = "|".join(sorted(_norm(h) for h in headers if h and h.strip()))
    return hashlib.sha1(canon.encode("utf-8")).hexdigest()[:16]


def detect_roles(headers: list[str], rows: list[dict]) -> dict:
    """Fuzzy role detection over headers + content.

    Returns {field_map, confidence, grouping, fingerprint, date_format}.
    field_map maps canonical role -> source column header.
    """
    field_map: dict[str, str] = {}
    scores: dict[str, float] = {}
    taken: set[str] = set()

    normed = {h: _norm(h) for h in headers}

    # Pass 1 — header synonyms (exact normalized match first, substring second)
    for role, syns in _ROLE_SYNONYMS.items():
        best = None
        best_score = 0.0
        for h in headers:
            if h in taken:
                continue
            n = normed[h]
            if not n:
                continue
            if n in syns:
                score = 1.0
            elif any(n == _norm(s) for s in syns):
                score = 1.0
            elif any(_norm(s) and _norm(s) in n for s in syns):
                score = 0.75
            else:
                continue
            if score > best_score:
                best, best_score = h, score
        if best:
            field_map[role] = best
            scores[role] = best_score
            taken.add(best)

    # Pass 2 — content sniffing
    def col_values(h):
        return [r.get(h) for r in rows]

    date_like = [
        h for h in headers
        if h not in taken and date_ratio(col_values(h)) >= 0.6
    ]
    for role in ("start_date", "end_date"):
        if role not in field_map and date_like:
            h = date_like.pop(0)
            field_map[role] = h
            scores[role] = 0.6
            taken.add(h)

    if "progress" not in field_map:
        for h in headers:
            if h in taken:
                continue
            vals = [v for v in col_values(h) if v is not None and str(v).strip() != ""]
            if not vals:
                continue
            nums = []
            for v in vals:
                try:
                    nums.append(float(str(v).replace("%", "").strip()))
                except ValueError:
                    nums = None
                    break
            if nums and all(0 <= n <= 100 for n in nums) and "%" in _norm(h) + "".join(
                str(v) for v in vals[:3]
            ):
                field_map["progress"] = h
                scores["progress"] = 0.6
                taken.add(h)
                break

    if "task_name" not in field_map:
        # Fall back: first mostly-unique text column
        for h in headers:
            if h in taken:
                continue
            vals = [str(v).strip() for v in col_values(h) if v is not None and str(v).strip()]
            if not vals or date_ratio(vals) >= 0.5:
                continue
            if len(set(vals)) >= max(1, int(0.8 * len(vals))):
                field_map["task_name"] = h
                scores["task_name"] = 0.5
                taken.add(h)
                break

    # Grouping suggestion — low-cardinality text column
    grouping: list[str] = []
    candidates = []
    n_rows = max(1, len(rows))
    for h in headers:
        if h == field_map.get("task_name"):
            continue
        if h in (field_map.get("start_date"), field_map.get("end_date"),
                 field_map.get("actual_start"), field_map.get("actual_end"),
                 field_map.get("progress")):
            continue
        vals = [str(v).strip() for v in col_values(h) if v is not None and str(v).strip()]
        if not vals or any(looks_like_date(v) for v in vals[:5]):
            continue
        card = len(set(vals))
        if 2 <= card <= min(12, max(2, n_rows // 2)) and len(vals) >= n_rows * 0.5:
            prio = next(
                (i for i, k in enumerate(_GROUPING_PRIORITY) if k in normed[h]),
                len(_GROUPING_PRIORITY),
            )
            candidates.append((prio, card, h))
    if candidates:
        candidates.sort()
        grouping = [candidates[0][2]]

    confidence = round(
        sum(scores.values()) / max(1, len(scores)), 2
    ) if scores else 0.0

    # Detected date wire format description
    date_format = "ISO"
    start_col = field_map.get("start_date")
    if start_col:
        vals = col_values(start_col)
        if column_dayfirst(vals):
            date_format = "D/M/Y"
        elif any(_SLASH_RE.match(str(v).strip()) for v in vals if v is not None):
            date_format = "M/D/Y"

    return {
        "field_map": field_map,
        "confidence": confidence,
        "grouping": grouping,
        "fingerprint": header_fingerprint(headers),
        "date_format": date_format,
    }


# ---------------------------------------------------------------------------
# Task building (rows -> hierarchical task dicts)
# ---------------------------------------------------------------------------


def iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    d = coerce_date(v)
    return d.isoformat() if d else None


def build_tasks(rows: list[dict], field_map: dict, grouping: list[str]) -> list[dict]:
    """Turn parsed rows into a flat task list with temp ids and group hierarchy.

    field_map maps canonical roles to source column headers. Accepts either
    'task_name' or legacy 'name' as the leaf-name role key.
    """
    name_col = field_map.get("task_name") or field_map.get("name")
    s_col = field_map.get("start_date")
    e_col = field_map.get("end_date")
    as_col = field_map.get("actual_start")
    ae_col = field_map.get("actual_end")
    p_col = field_map.get("progress")
    mapped = {c for c in (name_col, s_col, e_col, as_col, ae_col, p_col) if c}

    group_cache: dict[tuple, str] = {}
    tasks: list[dict] = []
    counter = {"n": 0}

    def _tid(prefix):
        counter["n"] += 1
        return f"{prefix}-{counter['n']}"

    for i, row in enumerate(rows):
        parent_temp = None
        path: list[str] = []
        for g_col in grouping or []:
            v = row.get(g_col)
            if v is None or str(v).strip() == "":
                continue
            path.append(str(v))
            key = tuple(path)
            if key not in group_cache:
                tid = _tid("g")
                group_cache[key] = tid
                tasks.append({
                    "temp_id": tid,
                    "parent_temp_id": parent_temp,
                    "task_name": str(v),
                    "kind": "group",
                    "sort_order": len(group_cache),
                    "fields": {g_col: str(v)},
                })
            parent_temp = group_cache[key]

        leaf_name = row.get(name_col) if name_col else None
        if leaf_name is None or str(leaf_name).strip() == "":
            continue

        progress = None
        if p_col and row.get(p_col) is not None:
            try:
                progress = int(round(float(str(row[p_col]).replace("%", "").strip())))
                progress = max(0, min(100, progress))
            except ValueError:
                progress = None

        fields = {}
        for col, val in row.items():
            if val is None or str(val).strip() == "":
                continue
            # Keep grouping columns in fields — rebucket needs them later.
            if col in mapped:
                continue
            if isinstance(val, (datetime, date)):
                val = iso(val)
            elif not isinstance(val, (str, int, float, bool)):
                val = str(val)
            fields[col] = val

        tasks.append({
            "temp_id": _tid("l"),
            "parent_temp_id": parent_temp,
            "task_name": str(leaf_name),
            "kind": "leaf",
            "start_date": iso(row.get(s_col)) if s_col else None,
            "end_date": iso(row.get(e_col)) if e_col else None,
            "actual_start": iso(row.get(as_col)) if as_col else None,
            "actual_end": iso(row.get(ae_col)) if ae_col else None,
            "progress": progress,
            "sort_order": i,
            "fields": fields,
        })
    return tasks


# ---------------------------------------------------------------------------
# Timeline-matrix detection (months across, modules down, cells = phase labels)
# ---------------------------------------------------------------------------

_MONTH_NUM = {abbr.lower(): i for i, abbr in enumerate(calendar.month_abbr) if abbr}
_MONTH_LABEL_RE = re.compile(r"^([A-Za-z]{3,9})\.?[\s\-/]+(\d{2}|\d{4})$")

#: minimum month-like columns for a row to qualify as a matrix header
_MIN_MONTH_COLS = 4
#: max distinct phase strings for the cell vocabulary to look "matrix-like"
_MAX_PHASE_VOCAB = 24


def parse_month_label(value) -> tuple[int, int] | None:
    """Parse 'Oct 24' / 'Oct-24' / 'Oct 2024' / 'October 24' -> (2024, 10).

    Also accepts date/datetime cells (XLSX month headers). Returns None when
    the value is not a month label. Two-digit years are 20xx ('Oct 24' = 2024).
    """
    if isinstance(value, (datetime, date)):
        return (value.year, value.month)
    s = str(value or "").strip()
    if not s:
        return None
    m = _MONTH_LABEL_RE.match(s)
    if not m:
        return None
    mon = m.group(1).lower()[:3]
    if mon not in _MONTH_NUM:
        return None
    y = int(m.group(2))
    if y < 100:
        y += 2000
    return (y, _MONTH_NUM[mon])


def month_bounds(year: int, month: int) -> tuple[str, str]:
    """First and last day of a month as ISO yyyy-mm-dd strings."""
    last = calendar.monthrange(year, month)[1]
    return f"{year:04d}-{month:02d}-01", f"{year:04d}-{month:02d}-{last:02d}"


def _cell(row: list, i: int) -> str:
    v = row[i] if i < len(row) else None
    return str(v).strip() if v is not None else ""


def detect_matrix(matrix: list[list]) -> dict | None:
    """Detect a timeline-matrix sheet in a raw cell matrix.

    Heuristics:
    - Within the first ~10 rows, find the row with the most month-like labels
      ("Oct 24", "Oct-24", "Oct 2024", real date cells). Needs >= 4 of them,
      chronologically non-decreasing. Handles a two-row header (quarters above
      months) naturally because the month row wins the count.
    - Columns before the first month column are label columns: the module-name
      column is the label column most filled in the body; an earlier sparse
      column (merged group numbers) is forward-filled as a group id.
    - Columns after the last month column (e.g. "Legends") are ignored.
    - Body cells under month columns must be non-date strings drawn from a
      small repeated vocabulary (phases).

    Returns {header_row, month_cols: [(col, (y, m)), ...], module_col,
    group_col, phases} or None when the sheet is not matrix-shaped.
    """
    if not matrix:
        return None

    best = None  # (count, row_idx, month_cols)
    for idx, row in enumerate(matrix[:10]):
        cols = []
        for i, v in enumerate(row or []):
            ym = parse_month_label(v)
            if ym:
                cols.append((i, ym))
        if len(cols) < _MIN_MONTH_COLS:
            continue
        # must be chronologically non-decreasing left to right
        keys = [y * 12 + m for _, (y, m) in cols]
        if any(b < a for a, b in zip(keys, keys[1:])):
            continue
        if best is None or len(cols) > best[0]:
            best = (len(cols), idx, cols)
    if not best:
        return None

    _, header_row, month_cols = best
    first_month_col = month_cols[0][0]
    if first_month_col == 0:
        return None  # no label columns to the left — not a matrix roadmap
    month_col_set = {c for c, _ in month_cols}

    body = matrix[header_row + 1 :]
    body = [r for r in body if r and any(_cell(r, i) for i in range(len(r)))]
    if len(body) < 2:
        return None

    # Label columns: pick the most-filled column left of the months as module
    fill = {
        i: sum(1 for r in body if _cell(r, i)) for i in range(first_month_col)
    }
    module_col = max(fill, key=lambda i: (fill[i], i))
    if fill[module_col] < max(2, len(body) // 2):
        return None
    group_col = None
    for i in range(first_month_col):
        if i != module_col and 0 < fill[i] < fill[module_col]:
            group_col = i
            break

    # Cell vocabulary check: month-column cells are short repeated strings,
    # not dates.
    values = []
    for r in body:
        for c in month_col_set:
            v = _cell(r, c)
            if v:
                values.append(v)
    if len(values) < 4:
        return None
    if sum(1 for v in values if looks_like_date(v)) > 0.2 * len(values):
        return None  # cells are dates — this is a table, not a phase matrix
    distinct = sorted(set(values))
    if len(distinct) > min(_MAX_PHASE_VOCAB, max(6, len(values) // 3)):
        return None

    return {
        "header_row": header_row,
        "month_cols": month_cols,
        "module_col": module_col,
        "group_col": group_col,
        "phases": distinct,
    }


def matrix_modules(matrix: list[list], info: dict) -> list[dict]:
    """Extract [{module, group, segments: [{phase, start, end}]}] from a
    detected matrix. Segments are contiguous same-phase month runs with ISO
    first/last-day dates. Phase strings are trimmed but otherwise kept as-is.
    """
    month_cols = info["month_cols"]
    module_col = info["module_col"]
    group_col = info.get("group_col")

    modules: list[dict] = []
    cur_group = ""
    for row in matrix[info["header_row"] + 1 :]:
        if not row or not any(_cell(row, i) for i in range(len(row))):
            continue
        if group_col is not None and _cell(row, group_col):
            cur_group = _cell(row, group_col)
        module = _cell(row, module_col)
        if not module:
            continue

        segments: list[dict] = []
        run_phase: str | None = None
        run_start: tuple[int, int] | None = None
        run_end: tuple[int, int] | None = None

        def _close():
            nonlocal run_phase, run_start, run_end
            if run_phase and run_start and run_end:
                start, _ = month_bounds(*run_start)
                _, end = month_bounds(*run_end)
                segments.append({"phase": run_phase, "start": start, "end": end})
            run_phase = run_start = run_end = None

        for col, ym in month_cols:
            val = _cell(row, col)
            if val and val == run_phase:
                run_end = ym
            else:
                _close()
                if val:
                    run_phase, run_start, run_end = val, ym, ym
        _close()

        modules.append({"module": module, "group": cur_group, "segments": segments})
    return modules


def build_matrix_tasks(matrix: list[list], info: dict | None = None) -> list[dict]:
    """Convert a timeline matrix to the flat task list bulk_insert_tasks expects.

    Each module row becomes a group; each contiguous same-phase month run in
    that row becomes a leaf named after the phase, spanning the first day of
    its first month to the last day of its last month, with
    fields {"Phase": ..., "Module": ...}.
    """
    if info is None:
        info = detect_matrix(matrix)
    if not info:
        raise ParserError("This file does not look like a timeline matrix")

    tasks: list[dict] = []
    counter = {"n": 0}

    def _tid(prefix):
        counter["n"] += 1
        return f"{prefix}-{counter['n']}"

    for m_idx, mod in enumerate(matrix_modules(matrix, info)):
        gid = _tid("g")
        fields = {"Module": mod["module"]}
        if mod["group"]:
            fields["Group"] = mod["group"]
        tasks.append({
            "temp_id": gid,
            "parent_temp_id": None,
            "task_name": mod["module"],
            "kind": "group",
            "sort_order": m_idx * 100,
            "fields": dict(fields),
        })
        for s_idx, seg in enumerate(mod["segments"]):
            tasks.append({
                "temp_id": _tid("l"),
                "parent_temp_id": gid,
                "task_name": seg["phase"],
                "kind": "leaf",
                "start_date": seg["start"],
                "end_date": seg["end"],
                "sort_order": m_idx * 100 + s_idx + 1,
                "fields": {**fields, "Phase": seg["phase"]},
            })
    return tasks


def matrix_summary(matrix: list[list], info: dict, preview_modules: int = 15) -> dict:
    """Summary + lightweight preview for the import wizard's matrix card."""
    modules = matrix_modules(matrix, info)
    (fy, fm) = info["month_cols"][0][1]
    (ly, lm) = info["month_cols"][-1][1]
    phases = sorted({s["phase"] for m in modules for s in m["segments"]})
    return {
        "modules_count": len(modules),
        "months_range": [month_bounds(fy, fm)[0], month_bounds(ly, lm)[1]],
        "phases": phases,
        "matrix_preview": [
            {"module": m["module"], "segments": m["segments"]}
            for m in modules[:preview_modules]
        ],
    }
